"""
services.sync_excel_data: Excel (data.xlsx) faylidagi ma'lumotlarni
bazaga (SQLite + JSON) xavfsiz, dublikatsiz va kadrlar tarixi bilan sinxronlash xizmati.
"""
import os
import re
import sys
import uuid
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
from collections import defaultdict
import openpyxl

from core.logger import logger
from database.data_manager import DataManager

# 1. Kirillcha-Lotincha o'girish jadvali
CYR2LAT = {
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ё': 'yo',
    'ж': 'j', 'з': 'z', 'и': 'i', 'й': 'y', 'к': 'k', 'л': 'l', 'м': 'm',
    'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'у': 'u',
    'ф': 'f', 'х': 'x', 'ҳ': 'x', 'ц': 'ts', 'ч': 'ch', 'ш': 'sh', 'щ': 'sh',
    'ъ': "'", 'ы': 'i', 'ь': '', 'э': 'e', 'ю': 'yu', 'я': 'ya',
    'ў': "o'", 'ғ': "g'", 'қ': 'q',
}

# Maxsus qisqartma va o'ziga xos yozilishlar lug'ati
SPECIAL_MAHALLA_MAP = {
    'а.навоий': 'Alisher Navoiy MFY',
    'а.темур': 'Amir Temur MFY',
    'ч.мурод': 'Cholik murod MFY',
    'м.улуғбек': 'Mirzo Ulugbek MFY',
    'а.яссавий': 'Ahmad Yassaviy MFY',
    'пахтакор': 'Paxtakor MFY',
    'анҳор': 'Anhor MFY',
    'муқимий': 'Muqumiy MFY',
    'навбаҳор': 'Navbahor MFY',
    'навруз': "Navro'z MFY",
    'обиҳаёт': 'Obihayot MFY',
    'сохибкор': 'Sohibkor MFY',
    'хазрати-боб': 'Hazrati bob MFY',
}

# Excel dagi lavozimlarni tizim toifalariga bog'lash
ROLE_TO_CATEGORY = {
    'МФЙ Раиси': 'Mahalla',
    'Ҳоким ёрдамчиси': 'Hokim yordamchisi',
    'Ёшлар етакчиси': 'Yoshlar yetakchisi',
    'Хотин-қизлар': 'Xotin qizlar',
    'Ижтимоий ҳимоя': 'Ijtimoiy xodim',
    'Инспектор': 'Profilaktika inspektori',
    'Солиқ': 'Soliq inspektori'
}

ROLE_TITLES = {
    'МФЙ Раиси': 'Mahalla Raisi',
    'Ҳоким ёрдамчиси': 'Hokim yordamchisi',
    'Ёшлар етакчиси': 'Yoshlar yetakchisi',
    'Хотин-қизлар': 'Xotin-qizlar faoli',
    'Ижтимоий ҳимоя': 'Ijtimoiy xodim (Inson)',
    'Инспектор': 'Profilaktika inspektori',
    'Солиқ': 'Soliq inspektori'
}

def translit_cyr_to_lat(text: str) -> str:
    """Kirill matnini Lotinchaga o'girish."""
    text = text.lower()
    res = [CYR2LAT.get(ch, ch) for ch in text]
    return "".join(res)

def clean_mahalla_key(text: str) -> str:
    """Solishtirish uchun mahallaning toza kaliti."""
    text = re.sub(r'\bmfy\b', '', text, flags=re.I)
    text = re.sub(r'[^a-z0-9]', '', text.lower())
    return text

def format_phone(phone_val: Any) -> str:
    """Telefon raqamini 'XX XXX XX XX' formatiga keltirish."""
    if not phone_val:
        return ""
    digits = re.sub(r'\D', '', str(phone_val))
    if digits.startswith('998') and len(digits) == 12:
        digits = digits[3:]
    if len(digits) == 9:
        return f"{digits[0:2]} {digits[2:5]} {digits[5:7]} {digits[7:9]}"
    return str(phone_val).strip()

def format_fio(fio_val: Any) -> str:
    """F.I.SH ni to'g'ri Title Case va o'g'li/qizi formatiga keltirish."""
    if not fio_val:
        return ""
    words = str(fio_val).strip().split()
    res = []
    for w in words:
        w_lower = w.lower()
        if w_lower in ("o'g'li", "o‘g‘li", "o`g`li", "og'li"):
            res.append("o'g'li")
        elif w_lower in ("qizi", "qizi."):
            res.append("qizi")
        else:
            if '-' in w:
                parts = w.split('-')
                sub_res = []
                for p in parts:
                    if p.lower() in ("o'g'li", "o‘g‘li", "qizi"):
                        sub_res.append(p.lower())
                    else:
                        sub_res.append(p.capitalize())
                res.append("-".join(sub_res))
            else:
                w_cap = w.capitalize()
                w_cap = re.sub(r"([GgOo])['‘`]([a-z])", lambda m: m.group(1).upper() + "'" + m.group(2).lower(), w_cap)
                res.append(w_cap)
    return " ".join(res)

class ExcelSyncService:
    """Excel faylidagi ma'lumotlarni bazaga xavfsiz kiritish klassi."""

    def __init__(self, data_manager: Optional[DataManager] = None):
        self.dm = data_manager or DataManager()

    def resolve_db_mahalla(self, xl_mahalla: str, db_mahallas: List[str]) -> Optional[str]:
        """Excel dagi mahalla nomini bazadagi aniq Lotincha MFY nomiga bog'lash."""
        xl_clean = xl_mahalla.strip().lower()
        if xl_clean in SPECIAL_MAHALLA_MAP:
            return SPECIAL_MAHALLA_MAP[xl_clean]

        c_xl = clean_mahalla_key(translit_cyr_to_lat(xl_mahalla))
        db_clean_map = {clean_mahalla_key(m): m for m in db_mahallas}

        if c_xl in db_clean_map:
            return db_clean_map[c_xl]

        for c_db, db_m in db_clean_map.items():
            if c_xl == c_db or c_xl in c_db or c_db in c_xl:
                return db_m

        return None

    def sync_from_excel(
        self,
        excel_path: str = "data.xlsx",
        include_inspector_and_soliq: bool = True
    ) -> Dict[str, Any]:
        """
        Excel faylidan barcha xodimlarni o'qib, bazani xavfsiz yangilash.
        Dublikatlar oldi olinadi, kadrlar tarixi qayd etiladi.
        """
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Excel fayli topilmadi: {excel_path}")

        logger.info(f"[EXCEL SYNC] Sinxronizatsiya boshlanmoqda: {excel_path}")

        # 1. Excel faylini o'qish
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet_name = 'Давомат' if 'Давомат' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # 2. Bazadagi barcha mahallalarni olish
        cursor = self.dm.sqlite.get_connection().cursor()
        cursor.execute("SELECT DISTINCT m FROM organizations WHERE s = 'Mahalla' ORDER BY m")
        db_mahallas = [r[0] for r in cursor.fetchall()]

        # 3. Excel qatorlarini yig'ish va dublikatlarni amaldagi ishchi bilan hal qilish
        rows_by_mahalla_role = defaultdict(list)
        for r in range(2, ws.max_row + 1):
            fio = ws.cell(r, 2).value
            m = ws.cell(r, 5).value
            role = ws.cell(r, 6).value
            status = ws.cell(r, 9).value
            reason = ws.cell(r, 8).value
            phone = ws.cell(r, 10).value

            if m and role and fio:
                rows_by_mahalla_role[(str(m).strip(), str(role).strip())].append({
                    'row': r,
                    'fio': format_fio(fio),
                    'phone': format_phone(phone),
                    'status': str(status) if status else '',
                    'reason': str(reason) if reason else ''
                })

        # Dublikatlarni tekshirish va faol xodimni saralash
        deduped_excel = {}
        for (m_name, r_name), entries in rows_by_mahalla_role.items():
            db_m = self.resolve_db_mahalla(m_name, db_mahallas)
            if not db_m:
                logger.warning(f"[EXCEL SYNC] Mahalla moslashtirilmadi: '{m_name}'")
                continue

            if len(entries) == 1:
                chosen = entries[0]
            else:
                # Ustuvorlik: amalda ishlayotgan xodim oldin turadi
                def get_priority(e):
                    st = e['status']
                    rs = e['reason']
                    if 'Ҳақиқатда ишланган' in st or 'вақтида келган' in st:
                        return 0
                    if 'Туғиш' in rs or 'таътил' in rs:
                        return 2
                    return 1

                sorted_entries = sorted(entries, key=get_priority)
                chosen = sorted_entries[0]

            deduped_excel[(db_m, r_name)] = chosen

        # 4. Hozirgi bazadagi ma'lumotlar xaritasi
        current_data = self.dm.data
        db_map = {}
        mahalla_inn_map = {}

        for item in current_data:
            s_val = item.get("s", "")
            m_val = item.get("m", "")
            inn_val = item.get("inn", "")
            db_map[(m_val, s_val)] = item
            if s_val == "Mahalla" and inn_val:
                mahalla_inn_map[m_val] = inn_val

        # 5. O'zgarishlarni qo'llash
        updated_count = 0
        unchanged_count = 0
        added_count = 0
        history_count = 0
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        for (db_m, r_name), x_item in deduped_excel.items():
            cat = ROLE_TO_CATEGORY.get(r_name)
            if not cat:
                continue

            if cat in ('Profilaktika inspektori', 'Soliq inspektori') and not include_inspector_and_soliq:
                continue

            new_fio = x_item['fio']
            new_phone = x_item['phone']

            if (db_m, cat) in db_map:
                # Mavjud yozuvni yangilash
                target_item = db_map[(db_m, cat)]
                old_fio = target_item.get("f", "")
                old_phone = target_item.get("t", "")

                has_changes = (old_fio != new_fio or old_phone != new_phone)
                if has_changes:
                    # Kadrlar tarixiga audit yozish
                    try:
                        self.dm.sqlite.add_staff_history(
                            org_id=target_item.get("id", ""),
                            mahalla=db_m,
                            role=ROLE_TITLES.get(r_name, cat),
                            full_name=new_fio,
                            phone=new_phone,
                            inn=target_item.get("inn", ""),
                            old_fio=old_fio,
                            new_fio=new_fio,
                            old_phone=old_phone,
                            new_phone=new_phone,
                            changed_by="ADMIN / EXCEL_SYNC",
                            reason="data.xlsx orqali yangilandi"
                        )
                        history_count += 1
                    except Exception as he:
                        logger.warning(f"Kadr tarixi yozishda ogohlantirish: {he}")

                    target_item["f"] = new_fio
                    target_item["t"] = new_phone
                    target_item["updated_at"] = now_str
                    updated_count += 1
                else:
                    unchanged_count += 1
            else:
                # Yangi yozuv qo'shish (yetishmayotgan xodimlar yoki yangi toifalar)
                m_inn = mahalla_inn_map.get(db_m, "")
                new_id = str(uuid.uuid4())
                new_org = {
                    "id": new_id,
                    "s": cat,
                    "m": db_m,
                    "f": new_fio,
                    "t": new_phone,
                    "inn": m_inn,
                    "izoh": ROLE_TITLES.get(r_name, cat),
                    "jshr": "",
                    "seriya": "",
                    "updated_at": now_str
                }
                current_data.append(new_org)
                db_map[(db_m, cat)] = new_org
                added_count += 1

                # Yangi xodim kadrlar tarixiga ham kiritiladi
                try:
                    self.dm.sqlite.add_staff_history(
                        org_id=new_id,
                        mahalla=db_m,
                        role=ROLE_TITLES.get(r_name, cat),
                        full_name=new_fio,
                        phone=new_phone,
                        inn=m_inn,
                        changed_by="ADMIN / EXCEL_SYNC",
                        reason="Yangi xodim bazaga kiritildi"
                    )
                    history_count += 1
                except Exception as he:
                    logger.warning(f"Yangi xodim tarixi yozishda ogohlantirish: {he}")

        # 6. Ma'lumotlarni Dual-Sync saqlash (JSON + SQLite)
        self.dm.data = current_data
        self.dm.save_data()

        summary = {
            "total_excel_rows": ws.max_row - 1,
            "deduped_entries": len(deduped_excel),
            "updated_staff": updated_count,
            "unchanged_staff": unchanged_count,
            "added_staff": added_count,
            "history_records_created": history_count,
            "total_orgs_in_db": len(self.dm.data)
        }

        logger.info(f"[EXCEL SYNC] Tugallandi: {summary}")
        return summary

def run_sync():
    """CLI orqali tezkor ishga tushirish funksiyasi."""
    service = ExcelSyncService()
    res = service.sync_from_excel("data.xlsx", include_inspector_and_soliq=True)
    print("Sinxronizatsiya natijasi:")
    for k, v in res.items():
        print(f"  {k}: {v}")

if __name__ == "__main__":
    run_sync()
