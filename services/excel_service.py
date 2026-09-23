from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill

def export_organizations_to_excel(data: List[Dict[str, Any]], filepath: str, category_name: str = "Barchasi") -> int:
    """Tashkilotlar ro'yxatini zamonaviy formatlangan Excel (.xlsx) fayliga yuklash."""
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Sheet nomi (maksimal 31 belgi)
    sheet_title = f"Tashkilotlar - {category_name}"[:31]
    ws.title = sheet_title

    # Sarlavhalar
    headers = ["Turi", "Nomi", "Rahbar", "Tel", "INN", "Izoh"]
    ws.append(headers)

    # Sarlavha dizayni
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Ma'lumotlarni yozish
    for item in data:
        ws.append([
            item.get("s") or "",
            item.get("m") or "",
            item.get("f") or "",
            item.get("t") or "",
            str(item.get("inn") or ""),
            item.get("izoh") or ""
        ])

    # Ustun kengliklarini moslash
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    wb.save(filepath)
    return len(data)

def import_organizations_from_file(filepath: str) -> tuple[List[Dict[str, Any]], List[str]]:
    """
    Excel (.xlsx, .xls) yoki CSV fayldan tashkilotlar ro'yxatini aqlli o'qish va validatsiya qilish.
    Qaytaradi: (yuklangan_tashkilotlar: list, ogohlantirishlar: list)
    """
    import os
    import csv
    import uuid
    from core.validators import sanitize_text, clean_inn, format_phone

    if not os.path.exists(filepath):
        return [], [f"Fayl topilmadi: {filepath}"]

    rows = []
    ext = os.path.splitext(filepath)[1].lower()

    if ext in [".xlsx", ".xlsm", ".xltx"]:
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if any(row):  # bo'sh bo'lmagan qatorlar
                    rows.append([str(c or "").strip() for c in row])
        except Exception as e:
            return [], [f"Excel faylini o'qishda xatolik: {e}"]
    elif ext == ".xls":
        try:
            import xlrd
            book = xlrd.open_workbook(filepath)
            sheet = book.sheet_by_index(0)
            for r in range(sheet.nrows):
                row_vals = sheet.row_values(r)
                if any(row_vals):
                    clean_row = []
                    for c in row_vals:
                        if isinstance(c, float) and c.is_integer():
                            clean_row.append(str(int(c)))
                        else:
                            clean_row.append(str(c or "").strip())
                    rows.append(clean_row)
        except Exception as e:
            return [], [f"Excel (.xls) faylini o'qishda xatolik: {e}"]
    elif ext == ".csv":
        try:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for r in reader:
                    if any(r):
                        rows.append([c.strip() for c in r])
        except Exception as e:
            return [], [f"CSV faylini o'qishda xatolik: {e}"]
    else:
        return [], ["Faqat Excel (.xlsx, .xls) yoki CSV (.csv) formatlari qo'llab-quvvatlanadi."]

    if not rows:
        return [], ["Fayl bo'sh."]

    # Sarlavhani topish va ustunlarni moslashtirish
    header = [h.lower() for h in rows[0]]
    col_map = {}

    for idx, h in enumerate(header):
        if any(k in h for k in ["turi", "toifa", "kategoriya", "soha", "sector", "type"]):
            col_map["s"] = idx
        elif any(k in h for k in ["nomi", "tashkilot", "mahalla", "mfy", "organization"]):
            col_map["m"] = idx
        elif any(k in h for k in ["rahbar", "f.i.sh", "fio", "xodim", "ism", "full_name"]):
            col_map["f"] = idx
        elif any(k in h for k in ["tel", "telefon", "phone", "mobil", "aloqa"]):
            col_map["t"] = idx
        elif any(k in h for k in ["inn", "stir", "soliq", "tax"]):
            col_map["inn"] = idx
        elif any(k in h for k in ["izoh", "qo'shimcha", "comment", "note"]):
            col_map["izoh"] = idx
        elif any(k in h for k in ["jshr", "jshshir", "pinfl"]):
            col_map["jshr"] = idx
        elif any(k in h for k in ["seriya", "pasport", "passport"]):
            col_map["seriya"] = idx

    # Standart zaxira ustun tartibi (agar sarlavha aniqlanmasa)
    if "m" not in col_map:
        # 0: turi, 1: nomi, 2: rahbar, 3: tel, 4: inn, 5: izoh
        col_map = {"s": 0, "m": 1, "f": 2, "t": 3, "inn": 4, "izoh": 5}
        start_row = 0
    else:
        start_row = 1

    imported_items: List[Dict[str, Any]] = []
    warnings: List[str] = []

    for r_idx, row in enumerate(rows[start_row:], start=start_row + 1):
        def get_val(key):
            idx = col_map.get(key)
            if idx is not None and idx < len(row):
                return row[idx].strip()
            return ""

        org_name = sanitize_text(get_val("m"))
        if not org_name:
            continue  # nomsiz qatorlar o'tkazib yuboriladi

        raw_inn = clean_inn(get_val("inn"))
        raw_tel = format_phone(get_val("t"))

        item = {
            "id": str(uuid.uuid4()),
            "s": sanitize_text(get_val("s")) or "Boshqa",
            "m": org_name,
            "f": sanitize_text(get_val("f")),
            "t": raw_tel,
            "inn": raw_inn,
            "izoh": sanitize_text(get_val("izoh")),
            "jshr": sanitize_text(get_val("jshr")),
            "seriya": sanitize_text(get_val("seriya"))
        }
        imported_items.append(item)

    if not imported_items:
        warnings.append("Fayldan yaroqli tashkilot ma'lumotlari topilmadi.")

    return imported_items, warnings

